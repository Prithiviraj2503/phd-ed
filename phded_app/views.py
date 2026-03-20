import pandas as pd

from django.shortcuts import render, redirect
from django.contrib.auth.models import User
from django.db import models
from django.http import JsonResponse, HttpResponse
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_http_methods
from django.views.decorators.csrf import ensure_csrf_cookie

from .models import UserProfile, Department, Course, CourseContent, Assignment, Question, StudentAssignmentAttempt, StudentAnswer, StudentSurvey
from .forms import LoginForm, CreateUserForm, BulkStudentUploadForm, CourseCreateForm, CourseContentForm, AssignmentCreateForm, StudentSurveyForm
from .utils import split_full_name, generate_password, send_welcome_email, send_course_content_notification, send_assignment_launch_notification
from .assignment_utils import extract_text_from_docx, generate_mcqs_from_text_groq, parse_excel_questions


def is_ajax(request):
    return request.headers.get('X-Requested-With') == 'XMLHttpRequest'


def login_required(f):
    def wrapper(request, *args, **kwargs):
        if not request.session.get('role'):
            return redirect('phded_app:login')
        return f(request, *args, **kwargs)
    return wrapper


def admin_required(f):
    def wrapper(request, *args, **kwargs):
        if request.session.get('role') != 'admin':
            return redirect('phded_app:login')
        return f(request, *args, **kwargs)
    return wrapper


def professor_required(f):
    def wrapper(request, *args, **kwargs):
        if request.session.get('role') != 'professor':
            return redirect('phded_app:login')
        return f(request, *args, **kwargs)
    return wrapper


def student_required(f):
    def wrapper(request, *args, **kwargs):
        if request.session.get('role') != 'student':
            return redirect('phded_app:login')
        return f(request, *args, **kwargs)
    return wrapper


def get_professor_user(request):
    uid = request.session.get('user_id')
    if not uid:
        return None
    return User.objects.filter(id=uid).first()


def get_student_user(request):
    uid = request.session.get('user_id')
    if not uid:
        return None
    return User.objects.filter(id=uid).first()


def get_course_students(course):
    return User.objects.filter(
        profile__role='student',
        profile__department=course.department.name,
    ).order_by('first_name', 'last_name', 'email')


def build_assignment_summary(assignment):
    students = list(get_course_students(assignment.course))
    attempts = list(
        StudentAssignmentAttempt.objects.filter(assignment=assignment)
        .select_related('student')
        .prefetch_related('answers__question')
        .order_by('student__first_name', 'student__last_name', 'student__email')
    )
    attempt_map = {attempt.student_id: attempt for attempt in attempts}
    submitted_attempts = [attempt for attempt in attempts if attempt.submitted_at]
    question_list = list(assignment.questions.all())

    summary_rows = []
    answer_rows = []
    for student in students:
        attempt = attempt_map.get(student.id)
        submitted = bool(attempt and attempt.submitted_at)
        score = attempt.weighted_score if submitted else 0
        max_score = attempt.max_weighted_score if attempt else sum(q.score_weight for q in question_list)
        summary_rows.append({
            'student': student,
            'attempt': attempt,
            'submitted': submitted,
            'score': score,
            'max_score': max_score,
        })
        if attempt:
            answers_by_question = {answer.question_id: answer for answer in attempt.answers.all()}
            for question in question_list:
                answer = answers_by_question.get(question.id)
                selected_answer = answer.selected_answer if answer else None
                is_correct = selected_answer == question.correct_answer if selected_answer else False
                answer_rows.append({
                    'student': student,
                    'question': question,
                    'selected_answer': selected_answer,
                    'selected_text': getattr(question, f'opt{selected_answer}', '') if selected_answer else '',
                    'correct_text': getattr(question, f'opt{question.correct_answer}', ''),
                    'is_correct': is_correct,
                    'submitted_at': attempt.submitted_at,
                })

    return {
        'students': students,
        'attempts': attempts,
        'submitted_count': len(submitted_attempts),
        'pending_count': max(len(students) - len(submitted_attempts), 0),
        'summary_rows': summary_rows,
        'answer_rows': answer_rows,
    }


def normalize_text(value):
    return str(value or '').strip()


def get_or_create_department_name(department_name):
    department_name = normalize_text(department_name)
    if not department_name:
        return ''
    existing = Department.objects.filter(name__iexact=department_name).first()
    if existing:
        return existing.name
    return Department.objects.create(name=department_name).name


def create_portal_user(*, role, full_name, email, phone='', department='', college='', address=''):
    email = normalize_text(email).lower()
    department_name = get_or_create_department_name(department)
    first_name, last_name = split_full_name(full_name)
    password = generate_password()

    user = User.objects.create_user(
        username=email,
        email=email,
        password=password,
        first_name=first_name,
        last_name=last_name,
    )
    UserProfile.objects.create(
        user=user,
        role=role,
        phone=normalize_text(phone),
        department=department_name,
        college=normalize_text(college),
        address=normalize_text(address),
    )
    email_sent = send_welcome_email(
        email=email,
        password=password,
        first_name=user.first_name or user.username,
        role=role,
    )
    return user, email_sent


def parse_student_upload_rows(uploaded_file):
    file_name = (uploaded_file.name or '').lower()
    if file_name.endswith('.csv'):
        dataframe = pd.read_csv(uploaded_file)
    elif file_name.endswith('.xlsx'):
        dataframe = pd.read_excel(uploaded_file)
    else:
        raise ValueError('Only .xlsx and .csv files are supported.')

    dataframe = dataframe.fillna('')
    rows = []
    for record in dataframe.to_dict(orient='records'):
        row_dict = {str(key): value for key, value in record.items()}
        row_dict['__row_values__'] = list(record.values())
        if any(normalize_text(value) for value in row_dict['__row_values__']):
            rows.append(row_dict)
    return rows


def map_student_upload_row(row):
    normalized = {
        normalize_text(key).lower().replace(' ', '').replace('_', '').replace('(optional)', ''): value
        for key, value in row.items()
        if key != '__row_values__'
    }
    mapped = {
        'full_name': normalize_text(
            normalized.get('studentname')
            or normalized.get('name')
            or normalized.get('fullname')
        ),
        'email': normalize_text(normalized.get('email')),
        'college': normalize_text(normalized.get('college')),
        'address': normalize_text(normalized.get('address')),
        'phone': normalize_text(normalized.get('phone')),
        'department': normalize_text(normalized.get('department')),
    }
    if mapped['full_name'] and mapped['email'] and mapped['department']:
        return mapped

    row_values = [normalize_text(value) for value in row.get('__row_values__', [])]
    if len(row_values) >= 7:
        mapped['full_name'] = mapped['full_name'] or row_values[1]
        mapped['email'] = mapped['email'] or row_values[2]
        mapped['college'] = mapped['college'] or row_values[3]
        mapped['address'] = mapped['address'] or row_values[4]
        mapped['phone'] = mapped['phone'] or row_values[5]
        mapped['department'] = mapped['department'] or row_values[6]
    elif len(row_values) >= 6:
        mapped['full_name'] = mapped['full_name'] or row_values[0]
        mapped['email'] = mapped['email'] or row_values[1]
        mapped['college'] = mapped['college'] or row_values[2]
        mapped['address'] = mapped['address'] or row_values[3]
        mapped['phone'] = mapped['phone'] or row_values[4]
        mapped['department'] = mapped['department'] or row_values[5]
    return mapped


def get_student_learning_summary(student):
    department_name = getattr(getattr(student, 'profile', None), 'department', '')
    courses = Course.objects.filter(department__name=department_name)
    total_courses = courses.count()
    total_contents = CourseContent.objects.filter(course__in=courses).count()
    assignment_attempts = StudentAssignmentAttempt.objects.filter(student=student)
    total_attempts = assignment_attempts.count()
    completed_attempts = assignment_attempts.filter(submitted_at__isnull=False).count()
    return {
        'department_name': department_name,
        'total_courses': total_courses,
        'total_contents': total_contents,
        'total_attempts': total_attempts,
        'completed_attempts': completed_attempts,
        'online_learning_activities': total_contents + total_attempts,
    }


def get_or_create_student_survey(student):
    survey, _ = StudentSurvey.objects.get_or_create(student=student)
    return survey


def get_visible_student_queryset(request):
    if request.session.get('role') == 'admin':
        return User.objects.filter(profile__role='student').select_related('profile').order_by('first_name', 'last_name', 'email')
    professor = get_professor_user(request)
    if not professor:
        return User.objects.none()
    department_names = Course.objects.filter(professor=professor).values_list('department__name', flat=True).distinct()
    return User.objects.filter(
        profile__role='student',
        profile__department__in=department_names,
    ).select_related('profile').order_by('first_name', 'last_name', 'email')


def build_student_survey_rows(request):
    rows = []
    for student in get_visible_student_queryset(request):
        survey = getattr(student, 'survey', None)
        learning = get_student_learning_summary(student)
        rows.append({
            'student': student,
            'profile': getattr(student, 'profile', None),
            'survey': survey,
            'learning': learning,
        })
    return rows


@ensure_csrf_cookie
@require_http_methods(["GET", "POST"])
def login_view(request):
    if request.method == 'GET':
        if request.session.get('role'):
            return redirect('phded_app:dashboard')
        form = LoginForm()
        return render(request, 'login.html', {'form': form})

    form = LoginForm(request.POST)
    if not form.is_valid():
        if is_ajax(request):
            return JsonResponse({
                'success': False,
                'errors': dict(form.errors),
                'error': 'Please fix the errors below.'
            }, status=400)
        return render(request, 'login.html', {'form': form, 'error_message': 'Please fix the errors below.'})

    email = form.cleaned_data['email'].strip().lower()
    password = form.cleaned_data['password']
    role = form.cleaned_data['role']

    if role == 'admin':
        try:
            user = User.objects.get(email__iexact=email)
            profile = UserProfile.objects.filter(user=user).first()
            is_admin_user = (
                (profile and profile.role == 'admin')
                or user.is_superuser
                or user.is_staff
            )
            if not is_admin_user:
                err = 'No admin account found for this email.'
                if is_ajax(request):
                    return JsonResponse({'success': False, 'error': err}, status=400)
                return render(request, 'login.html', {'form': form, 'error_message': err})
            if not user.check_password(password):
                err = 'Invalid admin credentials.'
                if is_ajax(request):
                    return JsonResponse({'success': False, 'error': err}, status=400)
                return render(request, 'login.html', {'form': form, 'error_message': err})

            request.session['user_id'] = user.id
            request.session['role'] = 'admin'
            request.session['user_name'] = user.get_full_name() or user.username or 'Admin'
            request.session.set_expiry(86400 * 7)
            if is_ajax(request):
                return JsonResponse({'success': True, 'redirect': '/dashboard/'})
            return redirect('phded_app:dashboard')
        except User.DoesNotExist:
            err = 'No admin account found for this email.'
            if is_ajax(request):
                return JsonResponse({'success': False, 'error': err}, status=400)
            return render(request, 'login.html', {'form': form, 'error_message': err})

    # Student or Professor: look up by email and check password
    try:
        user = User.objects.get(email__iexact=email)
        profile = user.profile
        if profile.role != role:
            err = 'No account found for this email and role.'
            if is_ajax(request):
                return JsonResponse({'success': False, 'error': err}, status=400)
            return render(request, 'login.html', {'form': form, 'error_message': err})
        if not user.check_password(password):
            err = 'Invalid password.'
            if is_ajax(request):
                return JsonResponse({'success': False, 'error': err}, status=400)
            return render(request, 'login.html', {'form': form, 'error_message': err})
        request.session['user_id'] = user.id
        request.session['role'] = profile.role
        request.session['user_name'] = user.get_full_name() or user.username
        request.session.set_expiry(86400 * 7)
        if is_ajax(request):
            return JsonResponse({'success': True, 'redirect': '/dashboard/'})
        return redirect('phded_app:dashboard')
    except (User.DoesNotExist, UserProfile.DoesNotExist):
        err = 'No account found for this email and role.'
        if is_ajax(request):
            return JsonResponse({'success': False, 'error': err}, status=400)
        return render(request, 'login.html', {'form': form, 'error_message': err})


@require_http_methods(["GET", "POST"])
def logout_view(request):
    request.session.flush()
    if is_ajax(request):
        return JsonResponse({'success': True, 'redirect': '/login/'})
    return redirect('phded_app:login')


@login_required
@require_http_methods(["GET", "POST"])
def dashboard_view(request):
    if request.session.get('role') == 'professor':
        return redirect('phded_app:professor_dashboard')
    if request.session.get('role') == 'student':
        return redirect('phded_app:student_dashboard')
    if request.session.get('role') != 'admin':
        ctx = {'user_name': request.session.get('user_name', 'User'), 'role': request.session.get('role', '')}
        if is_ajax(request):
            html = render(request, 'partials/dashboard_simple_content.html', ctx).content.decode()
            return JsonResponse({'html': html, 'breadcrumb_title': 'Welcome', 'breadcrumb_subtitle': 'You are logged in'})
        return render(request, 'dashboard_simple.html', ctx)
    students_count = UserProfile.objects.filter(role='student').count()
    professors_count = UserProfile.objects.filter(role='professor').count()
    total_users = UserProfile.objects.count()
    ctx = {
        'students_count': students_count,
        'professors_count': professors_count,
        'total_users': total_users,
    }
    if is_ajax(request):
        html = render(request, 'partials/dashboard_content.html', ctx).content.decode()
        return JsonResponse({'html': html, 'breadcrumb_title': 'Dashboard', 'breadcrumb_subtitle': 'Manage students and professors'})
    return render(request, 'dashboard.html', ctx)


@admin_required
@require_http_methods(["GET", "POST"])
def create_user_view(request):
    role = request.GET.get('role') or request.POST.get('role', 'student')
    if role not in ('student', 'professor'):
        if is_ajax(request):
            return JsonResponse({'success': False, 'error': 'Invalid role'}, status=400)
        return redirect('phded_app:dashboard')
    page_title = f'Add {"Professor" if role == "professor" else "Student"}'
    user_role = role

    if request.method == 'GET':
        form = CreateUserForm(role=user_role)
        bulk_form = BulkStudentUploadForm() if user_role == 'student' else None
        ctx = {'form': form, 'bulk_form': bulk_form, 'page_title': page_title, 'user_role': user_role}
        if is_ajax(request):
            html = render(request, 'partials/create_user_content.html', ctx).content.decode()
            return JsonResponse({'html': html, 'breadcrumb_title': page_title, 'breadcrumb_subtitle': f'Create new {user_role}'})
        return render(request, 'create_user.html', ctx)

    form = CreateUserForm(request.POST, role=user_role)
    if not form.is_valid():
        if is_ajax(request):
            return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
        return render(request, 'create_user.html', {'form': form, 'bulk_form': BulkStudentUploadForm() if user_role == 'student' else None, 'page_title': page_title, 'user_role': user_role})

    email = form.cleaned_data['email'].strip().lower()
    if User.objects.filter(email__iexact=email).exists():
        form.add_error('email', 'A user with this email already exists.')
        if is_ajax(request):
            return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
        return render(request, 'create_user.html', {'form': form, 'bulk_form': BulkStudentUploadForm() if user_role == 'student' else None, 'page_title': page_title, 'user_role': user_role})

    user, email_sent = create_portal_user(
        role=role,
        full_name=form.cleaned_data['full_name'],
        email=email,
        phone=form.cleaned_data.get('phone') or '',
        department=form.cleaned_data.get('department') or '',
        college=form.cleaned_data.get('college') or '',
        address=form.cleaned_data.get('address') or '',
    )
    if is_ajax(request):
        return JsonResponse({
            'success': True,
            'message': f'{user_role.title()} created. {"Welcome email sent" if email_sent else "Account created, but welcome email could not be sent"} to {email}.',
            'redirect': '/users/',
        })
    return redirect('phded_app:user_list')


@admin_required
@require_http_methods(["POST"])
def bulk_student_upload_view(request):
    form = BulkStudentUploadForm(request.POST, request.FILES)
    if not form.is_valid():
        return JsonResponse({'success': False, 'errors': dict(form.errors), 'error': 'Please upload a valid student sheet.'}, status=400)

    try:
        rows = parse_student_upload_rows(form.cleaned_data['file'])
    except ValueError as exc:
        return JsonResponse({'success': False, 'errors': {'file': [str(exc)]}, 'error': str(exc)}, status=400)
    except Exception:
        return JsonResponse({'success': False, 'errors': {'file': ['The file could not be read.']}, 'error': 'The file could not be read.'}, status=400)

    created_count = 0
    skipped_count = 0
    emailed_count = 0
    row_errors = []

    for index, raw_row in enumerate(rows, start=2):
        mapped_row = map_student_upload_row(raw_row)
        full_name = mapped_row['full_name']
        email = mapped_row['email'].lower()
        department = mapped_row['department']

        if not full_name or not email or not department:
            row_errors.append(f'Row {index}: Student Name, Email, and Department are required.')
            continue
        if User.objects.filter(email__iexact=email).exists():
            skipped_count += 1
            continue

        _, email_sent = create_portal_user(
            role='student',
            full_name=full_name,
            email=email,
            phone=mapped_row['phone'],
            department=department,
            college=mapped_row['college'],
            address=mapped_row['address'],
        )
        created_count += 1
        if email_sent:
            emailed_count += 1

    message = f'Bulk upload finished. Created {created_count} student(s), skipped {skipped_count} existing record(s), sent {emailed_count} welcome email(s).'
    if row_errors:
        message = f'{message} {len(row_errors)} row(s) had missing required fields.'

    response = {
        'success': True,
        'message': message,
        'redirect': '/users/',
    }
    if row_errors:
        response['message'] = f'{message} First issue: {row_errors[0]}'
    return JsonResponse(response)


@admin_required
@require_http_methods(["GET"])
def user_list_view(request):
    users = UserProfile.objects.select_related('user').all().order_by('-created_at')
    if is_ajax(request):
        html = render(request, 'partials/user_list_content.html', {'users': users}).content.decode()
        return JsonResponse({'html': html, 'breadcrumb_title': 'Users', 'breadcrumb_subtitle': 'Students and professors', 'hide_breadcrumb': True})
    return render(request, 'user_list.html', {'users': users})


@admin_required
@require_http_methods(["POST"])
def delete_user_view(request, user_id):
    try:
        user_profile = UserProfile.objects.get(user_id=user_id)
        user = user_profile.user
        username = user.get_full_name() or user.username
        user_profile.delete()
        user.delete()
        return JsonResponse({'success': True, 'message': f'User {username} deleted successfully'})
    except UserProfile.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'User not found'}, status=404)
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


# ---------- Admin: Courses ----------
@admin_required
@require_http_methods(["GET", "POST"])
def admin_course_list_view(request):
    courses = Course.objects.select_related('department', 'professor').order_by('-created_at')
    if request.method == 'POST' and is_ajax(request):
        return JsonResponse({'success': False, 'error': 'Use GET'})
    if is_ajax(request):
        html = render(request, 'partials/admin_course_list_content.html', {'courses': courses}).content.decode()
        return JsonResponse({'html': html, 'breadcrumb_title': 'Courses', 'breadcrumb_subtitle': 'Manage courses'})
    return render(request, 'admin_course_list.html', {'courses': courses})


@admin_required
@require_http_methods(["GET", "POST"])
def admin_course_create_view(request):
    if request.method == 'GET':
        form = CourseCreateForm()
        ctx = {'form': form}
        if is_ajax(request):
            html = render(request, 'partials/admin_course_create_content.html', ctx).content.decode()
            return JsonResponse({'html': html, 'breadcrumb_title': 'Create course', 'breadcrumb_subtitle': 'Assign department and professor'})
        return render(request, 'admin_course_create.html', ctx)
    form = CourseCreateForm(request.POST)
    if not form.is_valid():
        if is_ajax(request):
            return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
        return render(request, 'admin_course_create.html', {'form': form})
    Course.objects.create(
        name=form.cleaned_data['name'],
        code=form.cleaned_data.get('code') or '',
        department=form.cleaned_data['department'],
        professor=form.cleaned_data['professor'],
    )
    if is_ajax(request):
        return JsonResponse({'success': True, 'redirect': reverse('phded_app:admin_course_list')})
    return redirect('phded_app:admin_course_list')


# ---------- Professor: My courses, content, assignments ----------
@professor_required
@require_http_methods(["GET"])
def professor_dashboard_view(request):
    user = get_professor_user(request)
    if not user:
        return redirect('phded_app:login')
    courses = Course.objects.filter(professor=user).select_related('department').order_by('-created_at')
    if is_ajax(request):
        html = render(request, 'partials/professor_dashboard_content.html', {'courses': courses}).content.decode()
        return JsonResponse({'html': html, 'breadcrumb_title': 'My courses', 'breadcrumb_subtitle': 'Courses you teach', 'hide_breadcrumb': True})
    return render(request, 'professor_dashboard.html', {'courses': courses})


@login_required
@require_http_methods(["GET"])
def student_survey_list_view(request):
    if request.session.get('role') not in ('admin', 'professor'):
        return redirect('phded_app:dashboard')
    survey_rows = build_student_survey_rows(request)
    title = 'Student Surveys'
    subtitle = 'Academic, behaviour and demographic details'
    ctx = {
        'survey_rows': survey_rows,
        'page_title': title,
        'page_subtitle': subtitle,
        'user_role': request.session.get('role'),
    }
    if is_ajax(request):
        html = render(request, 'partials/student_survey_list_content.html', ctx).content.decode()
        return JsonResponse({'html': html, 'breadcrumb_title': title, 'breadcrumb_subtitle': subtitle})
    return render(request, 'student_survey_list.html', ctx)


@login_required
@require_http_methods(["GET"])
def student_survey_detail_view(request, student_id):
    if request.session.get('role') not in ('admin', 'professor'):
        return JsonResponse({'success': False, 'error': 'Unauthorized'}, status=403)
    student = get_visible_student_queryset(request).filter(id=student_id).first()
    if not student:
        return JsonResponse({'success': False, 'error': 'Student not found.'}, status=404)
    survey = getattr(student, 'survey', None)
    learning = get_student_learning_summary(student)
    html = render(request, 'partials/student_survey_detail_modal.html', {
        'student': student,
        'profile': getattr(student, 'profile', None),
        'survey': survey,
        'learning': learning,
    }).content.decode()
    return JsonResponse({'html': html})


@login_required
@require_http_methods(["GET"])
def student_survey_export_view(request):
    if request.session.get('role') not in ('admin', 'professor'):
        return redirect('phded_app:dashboard')
    try:
        from openpyxl import Workbook
    except ModuleNotFoundError:
        return HttpResponse('openpyxl is not installed in this environment.', status=500)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Student Surveys'
    sheet.append([
        'Student Name', 'Email', 'Department', 'College', 'Phone', 'Address',
        '10th Score', '12th Score', 'Attendance Record', 'Academic Level', 'Course Completion Rate',
        'Study Hours Daily', 'Gaming Hours', 'Social Media Hours', 'Sleep Hours',
        'Extra Curricular Activities', 'Extra Curricular Hours',
        'DOB', 'Age', 'Gender', 'Online Learning Activities', 'Quiz Attempts', 'Completed Quizzes', 'Available Course Contents',
    ])
    for row in build_student_survey_rows(request):
        student = row['student']
        profile = row['profile']
        survey = row['survey']
        learning = row['learning']
        sheet.append([
            student.get_full_name() or student.username,
            student.email,
            getattr(profile, 'department', ''),
            getattr(profile, 'college', ''),
            getattr(profile, 'phone', ''),
            getattr(profile, 'address', ''),
            getattr(survey, 'tenth_score', ''),
            getattr(survey, 'twelfth_score', ''),
            getattr(survey, 'attendance_record', ''),
            survey.get_academic_level_display() if survey and survey.academic_level else '',
            getattr(survey, 'course_completion_rate', ''),
            getattr(survey, 'study_hours_daily', ''),
            getattr(survey, 'gaming_hours', ''),
            getattr(survey, 'social_media_hours', ''),
            getattr(survey, 'sleep_hours', ''),
            getattr(survey, 'extra_curricular_activities', ''),
            getattr(survey, 'extra_curricular_hours', ''),
            survey.dob.isoformat() if survey and survey.dob else '',
            survey.age if survey else '',
            survey.get_gender_display() if survey and survey.gender else '',
            learning['online_learning_activities'],
            learning['total_attempts'],
            learning['completed_attempts'],
            learning['total_contents'],
        ])
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="student_surveys.xlsx"'
    workbook.save(response)
    return response


@professor_required
@require_http_methods(["GET"])
def professor_course_detail_view(request, course_id):
    user = get_professor_user(request)
    course = Course.objects.filter(professor=user, id=course_id).select_related('department').prefetch_related('contents', 'assignments__questions').first()
    if not course:
        return redirect('phded_app:professor_dashboard')
    if is_ajax(request):
        html = render(request, 'partials/professor_course_detail_content.html', {'course': course}).content.decode()
        return JsonResponse({'html': html, 'breadcrumb_title': course.name, 'breadcrumb_subtitle': 'Contents & assignments', 'hide_breadcrumb': True})
    return render(request, 'professor_course_detail.html', {'course': course})


@professor_required
@require_http_methods(["GET", "POST"])
def professor_upload_content_view(request, course_id):
    user = get_professor_user(request)
    course = Course.objects.filter(professor=user, id=course_id).first()
    if not course:
        return redirect('phded_app:professor_dashboard')
    if request.method == 'GET':
        form = CourseContentForm()
        ctx = {'form': form, 'course': course}
        if is_ajax(request):
            html = render(request, 'partials/professor_upload_content_content.html', ctx).content.decode()
            return JsonResponse({'html': html, 'breadcrumb_title': 'Upload content', 'breadcrumb_subtitle': course.name})
        return render(request, 'professor_upload_content.html', ctx)
    form = CourseContentForm(request.POST, request.FILES)
    if not form.is_valid():
        if is_ajax(request):
            return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
        return render(request, 'professor_upload_content.html', {'form': form, 'course': course})
    uploaded_file = request.FILES.get('file')
    if uploaded_file and not uploaded_file.name.lower().endswith('.docx'):
        form.add_error('file', 'Only Word .docx files can be uploaded for course content.')
        if is_ajax(request):
            return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
        return render(request, 'professor_upload_content.html', {'form': form, 'course': course})
    max_order = course.contents.aggregate(models.Max('order'))['order__max'] or 0
    content = CourseContent.objects.create(
        course=course,
        title=form.cleaned_data['title'],
        file=uploaded_file or None,
        order=max_order + 1,
    )
    notified_count = send_course_content_notification(course, user, content)
    if is_ajax(request):
        return JsonResponse({
            'success': True,
            'message': f'Content uploaded successfully. {notified_count} student(s) notified.',
            'redirect': f'/professor/course/{course_id}/'
        })
    return redirect('phded_app:professor_course_detail', course_id=course_id)


@professor_required
@require_http_methods(["GET", "POST"])
def professor_assignment_create_view(request, course_id):
    user = get_professor_user(request)
    course = Course.objects.filter(professor=user, id=course_id).first()
    if not course:
        return redirect('phded_app:professor_dashboard')
    if request.method == 'GET':
        form = AssignmentCreateForm(course=course)
        ctx = {'form': form, 'course': course}
        if is_ajax(request):
            html = render(request, 'partials/professor_assignment_create_content.html', ctx).content.decode()
            return JsonResponse({'html': html, 'breadcrumb_title': 'Create assignment', 'breadcrumb_subtitle': course.name})
        return render(request, 'professor_assignment_create.html', ctx)
    form = AssignmentCreateForm(request.POST, request.FILES, course=course)
    if not form.is_valid():
        if is_ajax(request):
            return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
        return render(request, 'professor_assignment_create.html', {'form': form, 'course': course})
    title = form.cleaned_data['title']
    source = form.cleaned_data['source_type']
    if source == 'docx' and not form.cleaned_data.get('content_source'):
        form.add_error('content_source', 'Please select uploaded course content.')
        if is_ajax(request):
            return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
        return render(request, 'professor_assignment_create.html', {'form': form, 'course': course})
    if source == 'excel' and not request.FILES.get('excel_file'):
        form.add_error('excel_file', 'Please upload an Excel file.')
        if is_ajax(request):
            return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
        return render(request, 'professor_assignment_create.html', {'form': form, 'course': course})
    assignment = Assignment.objects.create(course=course, title=title, created_by=user)
    questions_data = []
    generation_error = None
    if source == 'docx' and form.cleaned_data.get('content_source'):
        selected_content = form.cleaned_data['content_source']
        if selected_content.course_id != course.id:
            form.add_error('content_source', 'Invalid course content selected.')
            if is_ajax(request):
                return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
            return render(request, 'professor_assignment_create.html', {'form': form, 'course': course})
        if not selected_content.file:
            form.add_error('content_source', 'The selected course content does not have a file.')
            if is_ajax(request):
                return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
            return render(request, 'professor_assignment_create.html', {'form': form, 'course': course})
        file_name = selected_content.file.name.lower()
        if not file_name.endswith('.docx'):
            form.add_error('content_source', 'Please select uploaded course content with a Word .docx file.')
            if is_ajax(request):
                return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
            return render(request, 'professor_assignment_create.html', {'form': form, 'course': course})
        with selected_content.file.open('rb') as docx_file:
            text = extract_text_from_docx(docx_file)
        questions_data, generation_error = generate_mcqs_from_text_groq(text, num_questions=20)
    elif source == 'excel' and request.FILES.get('excel_file'):
        questions_data = parse_excel_questions(request.FILES['excel_file'])
    if not questions_data:
        assignment.delete()
        err = generation_error or 'No questions could be generated from the selected source.'
        if source == 'docx':
            form.add_error('content_source', err)
        else:
            form.add_error('excel_file', err)
        if is_ajax(request):
            return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
        return render(request, 'professor_assignment_create.html', {'form': form, 'course': course})
    for i, q in enumerate(questions_data):
        Question.objects.create(
            assignment=assignment,
            question_text=q['question_text'],
            opt1=q['opt1'], opt2=q['opt2'], opt3=q.get('opt3') or '', opt4=q.get('opt4') or '',
            correct_answer=q['correct_answer'],
            difficulty=q.get('difficulty', 'medium'),
            order=i + 1,
        )
    if is_ajax(request):
        return JsonResponse({'success': True, 'redirect': f'/professor/course/{course_id}/assignment/{assignment.id}/'})
    return redirect('phded_app:professor_assignment_detail', course_id=course_id, assignment_id=assignment.id)


@professor_required
@require_http_methods(["GET"])
def professor_assignment_detail_view(request, course_id, assignment_id):
    user = get_professor_user(request)
    course = Course.objects.filter(professor=user, id=course_id).first()
    if not course:
        return redirect('phded_app:professor_dashboard')
    assignment = Assignment.objects.filter(course=course, id=assignment_id).prefetch_related('questions').first()
    if not assignment:
        return redirect('phded_app:professor_course_detail', course_id=course_id)
    # Max possible weighted score for display
    total_weight = sum(q.score_weight for q in assignment.questions.all())
    ctx = {'course': course, 'assignment': assignment, 'total_weight': total_weight}
    if is_ajax(request):
        html = render(request, 'partials/professor_assignment_detail_content.html', ctx).content.decode()
        return JsonResponse({'html': html, 'breadcrumb_title': assignment.title, 'breadcrumb_subtitle': 'Questions & scoring'})
    return render(request, 'professor_assignment_detail.html', ctx)


@professor_required
@require_http_methods(["POST"])
def professor_assignment_launch_view(request, course_id, assignment_id):
    user = get_professor_user(request)
    course = Course.objects.filter(professor=user, id=course_id).first()
    if not course:
        return JsonResponse({'success': False, 'error': 'Course not found.'}, status=404)
    assignment = Assignment.objects.filter(course=course, id=assignment_id).prefetch_related('questions').first()
    if not assignment:
        return JsonResponse({'success': False, 'error': 'Assignment not found.'}, status=404)
    if assignment.is_launched:
        return JsonResponse({'success': False, 'error': 'Assignment is already launched.'}, status=400)
    if not assignment.questions.exists():
        return JsonResponse({'success': False, 'error': 'Assignment has no questions to launch.'}, status=400)

    assignment.launched_at = timezone.now()
    assignment.save(update_fields=['launched_at'])
    notified_count = send_assignment_launch_notification(course, assignment, user)
    return JsonResponse({
        'success': True,
        'message': f'Assignment launched successfully. {notified_count} student(s) notified.',
        'redirect': reverse('phded_app:professor_assignment_detail', kwargs={'course_id': course.id, 'assignment_id': assignment.id})
    })


@professor_required
@require_http_methods(["GET"])
def professor_assignment_summary_view(request, course_id, assignment_id):
    user = get_professor_user(request)
    course = Course.objects.filter(professor=user, id=course_id).first()
    if not course:
        return JsonResponse({'success': False, 'error': 'Course not found.'}, status=404)
    assignment = Assignment.objects.filter(course=course, id=assignment_id).prefetch_related('questions').first()
    if not assignment:
        return JsonResponse({'success': False, 'error': 'Assignment not found.'}, status=404)

    summary = build_assignment_summary(assignment)
    html = render(request, 'partials/professor_assignment_summary_modal.html', {
        'course': course,
        'assignment': assignment,
        **summary,
    }).content.decode()
    return JsonResponse({'html': html})


@professor_required
@require_http_methods(["GET"])
def professor_assignment_student_response_view(request, course_id, assignment_id, student_id):
    user = get_professor_user(request)
    course = Course.objects.filter(professor=user, id=course_id).first()
    if not course:
        return JsonResponse({'success': False, 'error': 'Course not found.'}, status=404)
    assignment = Assignment.objects.filter(course=course, id=assignment_id).prefetch_related('questions').first()
    if not assignment:
        return JsonResponse({'success': False, 'error': 'Assignment not found.'}, status=404)
    student = User.objects.filter(id=student_id, profile__role='student', profile__department=course.department.name).first()
    if not student:
        return JsonResponse({'success': False, 'error': 'Student not found.'}, status=404)

    summary = build_assignment_summary(assignment)
    selected_row = next((row for row in summary['summary_rows'] if row['student'].id == student.id), None)
    selected_answers = [row for row in summary['answer_rows'] if row['student'].id == student.id]
    html = render(request, 'partials/professor_student_response_modal.html', {
        'course': course,
        'assignment': assignment,
        'student': student,
        'summary_row': selected_row,
        'answer_rows': selected_answers,
    }).content.decode()
    return JsonResponse({'html': html})


@professor_required
@require_http_methods(["GET"])
def professor_assignment_summary_export_view(request, course_id, assignment_id):
    try:
        from openpyxl import Workbook
    except ModuleNotFoundError:
        return HttpResponse('openpyxl is not installed in this environment.', status=500)
    user = get_professor_user(request)
    course = Course.objects.filter(professor=user, id=course_id).first()
    if not course:
        return redirect('phded_app:professor_dashboard')
    assignment = Assignment.objects.filter(course=course, id=assignment_id).prefetch_related('questions').first()
    if not assignment:
        return redirect('phded_app:professor_course_detail', course_id=course_id)

    summary = build_assignment_summary(assignment)
    workbook = Workbook()
    ws_summary = workbook.active
    ws_summary.title = 'Summary'
    ws_summary.append(['Student Name', 'Email', 'Status', 'Score', 'Max Score', 'Submitted At'])
    for row in summary['summary_rows']:
        student = row['student']
        ws_summary.append([
            student.get_full_name() or student.username,
            student.email,
            'Taken' if row['submitted'] else 'Pending',
            row['score'],
            row['max_score'],
            row['attempt'].submitted_at.isoformat(sep=' ') if row['attempt'] and row['attempt'].submitted_at else '',
        ])

    ws_answers = workbook.create_sheet('Answer Tracking')
    ws_answers.append(['Student Name', 'Email', 'Question', 'Selected Answer', 'Correct Answer', 'Correct?', 'Difficulty', 'Weight', 'Submitted At'])
    for row in summary['answer_rows']:
        student = row['student']
        question = row['question']
        ws_answers.append([
            student.get_full_name() or student.username,
            student.email,
            question.question_text,
            row['selected_text'] or '',
            row['correct_text'],
            'Yes' if row['is_correct'] else 'No',
            question.difficulty,
            question.score_weight,
            row['submitted_at'].isoformat(sep=' ') if row['submitted_at'] else '',
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{assignment.title.replace(" ", "_")}_summary.xlsx"'
    workbook.save(response)
    return response


# ---------- Student: Courses by department, take assignment ----------
@student_required
@require_http_methods(["GET"])
def student_dashboard_view(request):
    user = get_student_user(request)
    if not user:
        return redirect('phded_app:login')
    profile = getattr(user, 'profile', None)
    dept_name = (profile and profile.department) or ''
    if not dept_name:
        courses = Course.objects.none()
    else:
        courses = Course.objects.filter(department__name=dept_name).select_related('department', 'professor').prefetch_related('assignments').order_by('-created_at')
    if is_ajax(request):
        html = render(request, 'partials/student_dashboard_content.html', {'courses': courses, 'dept_name': dept_name, 'survey_completed': hasattr(user, 'survey')}).content.decode()
        return JsonResponse({'html': html, 'breadcrumb_title': 'My courses', 'breadcrumb_subtitle': dept_name or 'Select your department', 'hide_breadcrumb': True})
    return render(request, 'student_dashboard.html', {'courses': courses, 'dept_name': dept_name, 'survey_completed': hasattr(user, 'survey')})


@student_required
@require_http_methods(["GET", "POST"])
def student_survey_view(request):
    student = get_student_user(request)
    if not student:
        return redirect('phded_app:login')
    survey = get_or_create_student_survey(student)
    if request.method == 'GET':
        form = StudentSurveyForm(instance=survey)
        learning = get_student_learning_summary(student)
        ctx = {
            'form': form,
            'survey': survey,
            'learning': learning,
            'student': student,
        }
        if is_ajax(request):
            html = render(request, 'partials/student_survey_form_content.html', ctx).content.decode()
            return JsonResponse({'html': html, 'breadcrumb_title': 'Student Survey', 'breadcrumb_subtitle': 'Save your academic, behaviour and demographic details'})
        return render(request, 'student_survey_form.html', ctx)

    form = StudentSurveyForm(request.POST, instance=survey)
    if not form.is_valid():
        if is_ajax(request):
            return JsonResponse({'success': False, 'errors': dict(form.errors)}, status=400)
        learning = get_student_learning_summary(student)
        return render(request, 'student_survey_form.html', {'form': form, 'survey': survey, 'learning': learning, 'student': student})

    form.save()
    if is_ajax(request):
        return JsonResponse({'success': True, 'message': 'Survey saved successfully.', 'redirect': reverse('phded_app:student_survey')})
    return redirect('phded_app:student_survey')


@student_required
@require_http_methods(["GET"])
def student_course_detail_view(request, course_id):
    user = get_student_user(request)
    if not user:
        return redirect('phded_app:login')
    profile = getattr(user, 'profile', None)
    dept_name = (profile and profile.department) or ''
    course = Course.objects.filter(id=course_id, department__name=dept_name).prefetch_related('assignments', 'contents').first()
    if not course:
        return redirect('phded_app:student_dashboard')
    completed_assignment_ids = set(
        StudentAssignmentAttempt.objects.filter(
            student=user, assignment__course=course, submitted_at__isnull=False
        ).values_list('assignment_id', flat=True)
    )
    launched_assignments = course.assignments.filter(launched_at__isnull=False)
    if is_ajax(request):
        html = render(request, 'partials/student_course_detail_content.html', {
            'course': course,
            'completed_assignment_ids': completed_assignment_ids,
            'launched_assignments': launched_assignments,
        }).content.decode()
        return JsonResponse({'html': html, 'breadcrumb_title': course.name, 'breadcrumb_subtitle': 'Assignments', 'hide_breadcrumb': True})
    return render(request, 'student_course_detail.html', {
        'course': course,
        'completed_assignment_ids': completed_assignment_ids,
        'launched_assignments': launched_assignments,
    })


@student_required
@require_http_methods(["GET", "POST"])
def student_assignment_take_view(request, assignment_id):
    user = get_student_user(request)
    if not user:
        return redirect('phded_app:login')
    profile = getattr(user, 'profile', None)
    dept_name = (profile and profile.department) or ''
    assignment = Assignment.objects.filter(id=assignment_id, course__department__name=dept_name, launched_at__isnull=False).prefetch_related('questions', 'course').first()
    if not assignment:
        return redirect('phded_app:student_dashboard')
    attempt, _ = StudentAssignmentAttempt.objects.get_or_create(student=user, assignment=assignment)
    if attempt.submitted_at:
        return redirect('phded_app:student_assignment_result', assignment_id=assignment_id)
    if request.method == 'POST':
        attempt.submitted_at = timezone.now()
        attempt.save()
        for q in assignment.questions.all():
            key = f'q_{q.id}'
            try:
                selected = int(request.POST.get(key, 0))
                if 1 <= selected <= 4:
                    StudentAnswer.objects.get_or_create(attempt=attempt, question=q, defaults={'selected_answer': selected})
            except (ValueError, TypeError):
                pass
        return redirect('phded_app:student_assignment_result', assignment_id=assignment_id)
    ctx = {'assignment': assignment, 'course': assignment.course}
    if is_ajax(request):
        html = render(request, 'partials/student_assignment_take_content.html', ctx).content.decode()
        return JsonResponse({'html': html, 'breadcrumb_title': assignment.title, 'breadcrumb_subtitle': 'Take test'})
    return render(request, 'student_assignment_take.html', ctx)


@student_required
@require_http_methods(["GET"])
def student_assignment_result_view(request, assignment_id):
    user = get_student_user(request)
    if not user:
        return redirect('phded_app:login')
    profile = getattr(user, 'profile', None)
    dept_name = (profile and profile.department) or ''
    assignment = Assignment.objects.filter(id=assignment_id, course__department__name=dept_name, launched_at__isnull=False).select_related('course').first()
    if not assignment:
        return redirect('phded_app:student_dashboard')
    attempt = StudentAssignmentAttempt.objects.filter(student=user, assignment=assignment).first()
    if not attempt or not attempt.submitted_at:
        return redirect('phded_app:student_assignment_take', assignment_id=assignment_id)
    score = attempt.weighted_score
    max_score = attempt.max_weighted_score
    percentage = round(100 * score / max_score, 1) if max_score else 0
    ctx = {'assignment': assignment, 'course': assignment.course, 'attempt': attempt, 'score': score, 'max_score': max_score, 'percentage': percentage}
    if is_ajax(request):
        html = render(request, 'partials/student_assignment_result_content.html', ctx).content.decode()
        return JsonResponse({'html': html, 'breadcrumb_title': assignment.title, 'breadcrumb_subtitle': 'Result'})
    return render(request, 'student_assignment_result.html', ctx)
